"""run sensor: notes with before/after, metrics, label sample + import, accuracy re-weighted to real totals."""
import csv
import json
from datetime import datetime, timedelta

from sqlalchemy import text

from app import sensor
from app.database.repository import DatabaseRepository
from tests.test_report_xlsx import TODAY, _seed


def _pieces(settings, n_ok=60, n_ng=60):
    repo = DatabaseRepository(settings.database_path)
    ts = (TODAY - timedelta(hours=2))
    with repo.engine.begin() as c:
        for i in range(n_ok + n_ng):
            ng = i >= n_ok
            tools = json.dumps([{"no": 1, "name": "AI", "status": "NG" if ng else "OK", "value": 40 if ng else 99}])
            c.execute(text(
                "INSERT INTO inspection (uid, inspection_id, timestamp, camera_id, analysis_status, score, tools_json,"
                " image_file, created_at, updated_at) VALUES (:u, :i, :t, 'IV4-01', :s, 90, :tj, :img, :c, :c)"),
                {"u": f"p{i}", "i": f"{i:05d}", "t": ts.strftime("%Y-%m-%d %H:%M:%S"), "s": "FAIL" if ng else "PASS",
                 "tj": tools, "img": f"p{i}.jpg", "c": datetime.utcnow()})
    repo.dispose()


def test_sample_labels_and_weighted_accuracy(settings, tmp_path):
    _seed(settings)
    _pieces(settings)
    repo = DatabaseRepository(settings.database_path)
    with repo.engine.begin() as c:                                   # the real totals of the hour the pieces came from
        c.execute(text(
            "INSERT INTO hourly_stats (hour, program_no, sensor_id, total, pass_count, fail_count, unknown_count,"
            " time_ms_sum, time_ms_count, time_ms_max, trigger_min, trigger_max, trigger_resets)"
            " VALUES (:h, 0, 'IV4-01', 12000, 11940, 60, 0, 1, 1, 1, 1, 12000, 0)"),
            {"h": (TODAY - timedelta(hours=2)).strftime("%Y-%m-%d %H")})
    path, n = sensor.make_sample(repo, 40, 7, TODAY, seed=1, base=tmp_path)
    assert n == 40
    rows = list(csv.DictReader(open(path, encoding="utf-8-sig")))
    assert sum(r["sensor_result"] == "NG" for r in rows) == 20 and all(r["truth"] == "" for r in rows)
    # a person checks: 2 of the 20 sensor-NG were really OK; 1 of the 20 sensor-OK was really NG
    fr = miss = 0
    for r in rows:
        r["truth"] = r["sensor_result"]
        if r["sensor_result"] == "NG" and fr < 2:
            r["truth"], fr = "OK", fr + 1
        elif r["sensor_result"] == "OK" and miss < 1:
            r["truth"], miss = "NG", miss + 1
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    saved, blank, problems = sensor.import_labels(repo, path, TODAY)
    assert saved == 40 and not problems
    acc = sensor.accuracy(repo)
    assert acc["cm"][("NG", "OK")] == 2 and acc["cm"][("OK", "NG")] == 1
    p_ok, p_ng = acc["pop"]
    want = 1 - (p_ok * 1 / 20 + p_ng * 2 / 20) / (p_ok + p_ng)
    assert abs(acc["weighted"] - want) < 1e-9 and acc["weighted"] > 0.9
    assert acc["fr_tools"]["T1"] == 2
    assert abs(acc["precision"] - 18 / 20) < 1e-9
    true_ng, missed = p_ng * 18 / 20, p_ok * 1 / 20
    assert abs(acc["recall"] - true_ng / (true_ng + missed)) < 1e-9 and 0 < acc["f1"] < 1
    out = sensor.export_accuracy(repo, tmp_path, TODAY)
    sensor.export_accuracy(repo, tmp_path, TODAY)                  # second run appends, header once
    got = list(csv.DictReader(open(out, encoding="utf-8-sig")))
    assert len(got) == 2 and got[0]["checked"] == "40" and got[0]["FP_sensorNG_truthOK"] == "2"
    assert float(got[0]["accuracy_pct"]) == round(100 * acc["weighted"], 2)
    assert float(got[0]["precision_pct"]) == 90.0 and float(got[0]["false_reject_pct"]) == 10.0
    repo.dispose()


def test_import_rejects_bad_rows(settings, tmp_path):
    _seed(settings)
    _pieces(settings, 2, 2)
    f = tmp_path / "l.csv"
    f.write_text("uid,truth\np0,OK\np1,maybe\nzzz,NG\np2,\n", encoding="utf-8")
    repo = DatabaseRepository(settings.database_path)
    saved, blank, problems = sensor.import_labels(repo, f, TODAY)
    repo.dispose()
    assert (saved, blank) == (1, 1) and len(problems) == 2
    f.write_text("a,b\n1,2\n", encoding="utf-8")
    repo = DatabaseRepository(settings.database_path)
    assert sensor.import_labels(repo, f, TODAY)[0] == 0
    repo.dispose()


def test_notes_show_before_after_and_report_runs(settings, tmp_path, capsys):
    _seed(settings)
    repo = DatabaseRepository(settings.database_path)
    sensor.add_note(repo, "Tool1 threshold 70->65", TODAY - timedelta(days=2) + timedelta(hours=12))
    lines = sensor.notes_section(repo, TODAY)
    repo.dispose()
    assert "Tool1 threshold 70->65" in lines[0]
    assert sensor.run(["--days", "5"], settings, base=tmp_path, now=TODAY) == 0
    out = capsys.readouterr().out
    assert "Lowest OK" in out and "no checked pieces yet" in out and "AI Differentiate" in out
    assert (tmp_path / "reports").exists()
    assert sensor.run(["--bogus"], settings, base=tmp_path) == 2


def test_sample_has_every_value_the_sensor_sends(settings, tmp_path):
    _seed(settings)
    _pieces(settings, 3, 3)
    repo = DatabaseRepository(settings.database_path)
    path, n = sensor.make_sample(repo, 6, 7, TODAY, seed=2, base=tmp_path)
    repo.dispose()
    rows = list(csv.DictReader(open(path, encoding="utf-8-sig")))
    assert {"program", "trigger_no", "time_ms", "tool1_status", "tool1_value", "folder", "truth"} <= set(rows[0])
    assert {r["tool1_value"] for r in rows} == {"40", "99"}


def test_label_workbook_has_image_links_and_imports_back(settings, tmp_path):
    from openpyxl import load_workbook
    _seed(settings)
    _pieces(settings, 3, 3)
    repo = DatabaseRepository(settings.database_path)
    path, n = sensor.make_sample(repo, 6, 7, TODAY, seed=3, base=tmp_path)
    book = sensor.write_label_workbook(path, r"D:\iv4\data\archive")
    wb = load_workbook(book)
    ws = wb.active
    assert ws["A2"].value.startswith('=HYPERLINK("D:\\iv4\\data\\archive\\') and ws["A2"].value.endswith('.jpg","open")')
    assert ws.data_validations.dataValidation[0].formula1 == '"OK,NG"'
    head = [c.value for c in ws[1]]
    col = head.index("truth") + 1
    for i in range(2, ws.max_row + 1):
        ws.cell(i, col, "NG" if ws.cell(i, head.index("sensor_result") + 1).value == "NG" else "OK")
    wb.save(book)
    saved, blank, problems = sensor.import_labels(repo, book, TODAY)
    repo.dispose()
    assert (saved, blank, problems) == (6, 0, [])
